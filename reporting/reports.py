#!/usr/bin/env python3
"""
ASTREX v3.0 — Report Generator
Генерация отчётов в форматах HTML, CSV, Excel, PDF
"""

import csv
import json
import io
import html as _html
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

from core.logging_setup import get_logger

log = get_logger("astrex.reporting")


# ═══════════════════════════════════════════════════════════════════════════════
# DATA MODELS
# ═══════════════════════════════════════════════════════════════════════════════

@dataclass
class ReportSection:
    """Секция отчёта"""
    title: str
    content: str = ""
    data: List[Dict[str, Any]] = field(default_factory=list)
    chart_type: Optional[str] = None  # bar, pie, timeline


@dataclass
class ReportData:
    """Данные для отчёта"""
    title: str = "ASTREX Report"
    subtitle: str = ""
    query: str = ""
    folder: str = ""
    generated_at: str = field(default_factory=lambda: datetime.now().isoformat())
    sections: List[ReportSection] = field(default_factory=list)
    results: List[Dict[str, Any]] = field(default_factory=list)
    stats: Dict[str, Any] = field(default_factory=dict)
    entities: Dict[str, List[str]] = field(default_factory=dict)
    graph_data: Optional[Dict] = None
    timeline: Optional[List[Dict]] = None

    @classmethod
    def from_scan_results(cls, query: str, folder: str, results: list,
                          stats: dict, entities: dict = None,
                          graph_data: dict = None, timeline: list = None):
        """Создать ReportData из результатов сканирования"""
        report = cls(
            title=f"ASTREX — Результаты поиска: {query}",
            subtitle=f"Папка: {folder}",
            query=query,
            folder=folder,
            stats=stats,
            entities=entities or {},
            graph_data=graph_data,
            timeline=timeline,
        )

        for r in results:
            d = r.to_dict() if hasattr(r, 'to_dict') else r
            report.results.append(d)

        # Summary section
        report.sections.append(ReportSection(
            title="Сводка",
            content=(
                f"Запрос: {query}\n"
                f"Найдено совпадений: {len(results)}\n"
                f"Всего файлов: {stats.get('total_files', '?')}\n"
                f"Время: {stats.get('duration_seconds', '?')}с"
            )
        ))

        # Entities section
        if entities:
            report.sections.append(ReportSection(
                title="Извлечённые сущности",
                data=[
                    {"type": k, "count": len(v), "values": ", ".join(v[:20])}
                    for k, v in entities.items() if v
                ]
            ))

        return report


# ═══════════════════════════════════════════════════════════════════════════════
# HTML REPORT
# ═══════════════════════════════════════════════════════════════════════════════

_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
  :root {{ --bg: #0a0e17; --fg: #c8d6e5; --accent: #00ff41;
           --card: #111827; --border: #1e3a5f; --red: #ff4444; }}
  * {{ margin:0; padding:0; box-sizing:border-box; }}
  body {{ background:var(--bg); color:var(--fg); font-family:'Courier New',monospace;
          padding:40px; line-height:1.6; }}
  h1 {{ color:var(--accent); font-size:28px; border-bottom:2px solid var(--accent);
        padding-bottom:10px; margin-bottom:20px; }}
  h2 {{ color:var(--accent); font-size:20px; margin:30px 0 15px; }}
  h3 {{ color:#4fc3f7; font-size:16px; margin:20px 0 10px; }}
  .meta {{ color:#666; font-size:13px; margin-bottom:30px; }}
  .stats {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(200px,1fr));
            gap:15px; margin:20px 0; }}
  .stat-card {{ background:var(--card); border:1px solid var(--border);
                border-radius:8px; padding:20px; text-align:center; }}
  .stat-value {{ font-size:32px; font-weight:bold; color:var(--accent); }}
  .stat-label {{ color:#888; font-size:12px; text-transform:uppercase; }}
  table {{ width:100%; border-collapse:collapse; margin:15px 0; }}
  th {{ background:#1a2744; color:var(--accent); text-align:left; padding:12px;
        font-size:13px; text-transform:uppercase; }}
  td {{ padding:10px 12px; border-bottom:1px solid var(--border); font-size:13px; }}
  tr:hover {{ background:#0d1b2a; }}
  .score {{ color:var(--accent); font-weight:bold; }}
  .entity-tag {{ display:inline-block; background:#1a3a2a; color:#4caf50;
                 padding:2px 8px; border-radius:4px; margin:2px; font-size:11px; }}
  .snippet {{ color:#999; font-size:12px; max-width:500px;
              overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }}
  .timeline {{ position:relative; margin:20px 0; padding-left:30px;
               border-left:2px solid var(--accent); }}
  .timeline-item {{ margin:15px 0; padding:10px 15px; background:var(--card);
                    border-radius:6px; position:relative; }}
  .timeline-item::before {{ content:''; position:absolute; left:-37px; top:15px;
                            width:12px; height:12px; background:var(--accent);
                            border-radius:50%; }}
  .timeline-date {{ color:var(--accent); font-weight:bold; }}
  .footer {{ margin-top:50px; padding-top:20px; border-top:1px solid var(--border);
             color:#555; font-size:11px; text-align:center; }}
</style>
</head>
<body>
<h1>{title}</h1>
<div class="meta">{subtitle}<br>Сгенерировано: {generated_at}</div>
{stats_html}
{sections_html}
{results_html}
{timeline_html}
<div class="footer">ASTREX v3.0 Intelligence System</div>
</body></html>"""


def _render_stats_html(stats: Dict) -> str:
    if not stats:
        return ""
    cards = []
    labels = {
        'total_files': 'Файлов',
        'matched_files': 'Совпадений',
        'processed_files': 'Обработано',
        'errors': 'Ошибок',
        'duration_seconds': 'Время (сек)',
        'files_per_second': 'Файлов/сек',
    }
    for key, label in labels.items():
        val = stats.get(key)
        if val is not None:
            if isinstance(val, float):
                val = f"{val:.1f}"
            cards.append(
                f'<div class="stat-card"><div class="stat-value">{val}</div>'
                f'<div class="stat-label">{label}</div></div>'
            )
    return f'<h2>Статистика</h2><div class="stats">{"".join(cards)}</div>'


def _render_sections_html(sections: List[ReportSection]) -> str:
    parts = []
    for sec in sections:
        parts.append(f"<h2>{_html.escape(str(sec.title))}</h2>")
        if sec.content:
            parts.append(f"<pre>{_html.escape(str(sec.content))}</pre>")
        if sec.data and len(sec.data) > 0 and isinstance(sec.data[0], dict):
            parts.append("<table><tr>")
            keys = list(sec.data[0].keys())
            for k in keys:
                parts.append(f"<th>{_html.escape(str(k))}</th>")
            parts.append("</tr>")
            for row in sec.data:
                parts.append("<tr>")
                for k in keys:
                    parts.append(f"<td>{_html.escape(str(row.get(k, '')))}</td>")
                parts.append("</tr>")
            parts.append("</table>")
    return "".join(parts)


def _render_results_html(results: List[Dict]) -> str:
    if not results:
        return ""
    parts = ['<h2>Результаты</h2><table>',
             '<tr><th>#</th><th>Файл</th><th>Score</th>'
             '<th>Сущности</th><th>Фрагмент</th></tr>']
    for i, r in enumerate(results, 1):
        entities_html = ""
        ents = r.get('entities', {})
        if isinstance(ents, dict):
            for etype, vals in ents.items():
                if isinstance(vals, list):
                    for v in vals[:5]:
                        entities_html += (
                            f'<span class="entity-tag">'
                            f'{_html.escape(str(etype))}: {_html.escape(str(v))}'
                            f'</span> '
                        )
        snippet = _html.escape((r.get('snippet', '') or '')[:200])
        score = r.get('score', 0)
        parts.append(
            f'<tr><td>{i}</td><td>{_html.escape(str(r.get("filename", "?")))}</td>'
            f'<td class="score">{score:.2f}</td>'
            f'<td>{entities_html}</td>'
            f'<td class="snippet">{snippet}</td></tr>'
        )
    parts.append("</table>")
    return "".join(parts)


def _render_timeline_html(timeline: Optional[List[Dict]]) -> str:
    if not timeline:
        return ""
    parts = ['<h2>Хронология</h2><div class="timeline">']
    for item in timeline[:100]:
        date = _html.escape(str(item.get('date', '?')))
        text = _html.escape(str(item.get('text', '')))
        source = _html.escape(str(item.get('source', '')))
        parts.append(
            f'<div class="timeline-item">'
            f'<div class="timeline-date">{date}</div>'
            f'<div>{text}</div>'
            f'<div style="color:#555;font-size:11px">{source}</div></div>'
        )
    parts.append("</div>")
    return "".join(parts)


def generate_html_report(data: ReportData, output_path: str) -> bool:
    """Генерация HTML отчёта"""
    try:
        html = _HTML_TEMPLATE.format(
            title=_html.escape(data.title),
            subtitle=_html.escape(data.subtitle),
            generated_at=_html.escape(data.generated_at),
            stats_html=_render_stats_html(data.stats),
            sections_html=_render_sections_html(data.sections),
            results_html=_render_results_html(data.results),
            timeline_html=_render_timeline_html(data.timeline),
        )
        Path(output_path).write_text(html, encoding='utf-8')
        log.info(f"HTML report saved: {output_path}")
        return True
    except Exception as e:
        log.error(f"HTML report failed: {e}")
        return False


# ═══════════════════════════════════════════════════════════════════════════════
# CSV REPORT
# ═══════════════════════════════════════════════════════════════════════════════

def generate_csv_report(data: ReportData, output_path: str) -> bool:
    """Генерация CSV отчёта"""
    try:
        with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([
                '№', 'Файл', 'Путь', 'Score',
                'Сущности', 'Фрагмент', 'Ошибка'
            ])
            for i, r in enumerate(data.results, 1):
                ents = r.get('entities', {})
                ents_str = "; ".join(
                    f"{k}: {', '.join(v[:5])}"
                    for k, v in ents.items() if isinstance(v, list) and v
                ) if isinstance(ents, dict) else ""
                writer.writerow([
                    i,
                    r.get('filename', ''),
                    r.get('path', ''),
                    f"{r.get('score', 0):.3f}",
                    ents_str,
                    (r.get('snippet', '') or '')[:500],
                    r.get('error', ''),
                ])
        log.info(f"CSV report saved: {output_path}")
        return True
    except Exception as e:
        log.error(f"CSV report failed: {e}")
        return False


# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL REPORT
# ═══════════════════════════════════════════════════════════════════════════════

def generate_excel_report(data: ReportData, output_path: str) -> bool:
    """Генерация Excel отчёта (openpyxl)"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        log.error("openpyxl not installed. Run: pip install openpyxl")
        return False

    try:
        wb = Workbook()

        # --- Sheet 1: Results ---
        ws = wb.active
        ws.title = "Результаты"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1a2744", end_color="1a2744",
                                  fill_type="solid")

        headers = ['№', 'Файл', 'Путь', 'Score', 'Сущности', 'Фрагмент']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill

        for i, r in enumerate(data.results, 1):
            ents = r.get('entities', {})
            ents_str = "; ".join(
                f"{k}: {', '.join(v[:5])}"
                for k, v in ents.items() if isinstance(v, list) and v
            ) if isinstance(ents, dict) else ""
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=r.get('filename', ''))
            ws.cell(row=i + 1, column=3, value=r.get('path', ''))
            ws.cell(row=i + 1, column=4, value=round(r.get('score', 0), 3))
            ws.cell(row=i + 1, column=5, value=ents_str)
            ws.cell(row=i + 1, column=6, value=(r.get('snippet', '') or '')[:500])

        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['E'].width = 40
        ws.column_dimensions['F'].width = 60

        # --- Sheet 2: Entities ---
        if data.entities:
            ws2 = wb.create_sheet("Сущности")
            for col, h in enumerate(['Тип', 'Количество', 'Значения'], 1):
                cell = ws2.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
            row = 2
            for etype, vals in data.entities.items():
                if vals:
                    ws2.cell(row=row, column=1, value=etype)
                    ws2.cell(row=row, column=2, value=len(vals))
                    ws2.cell(row=row, column=3, value=", ".join(vals[:30]))
                    row += 1

        # --- Sheet 3: Stats ---
        ws3 = wb.create_sheet("Статистика")
        for col, h in enumerate(['Параметр', 'Значение'], 1):
            cell = ws3.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
        row = 2
        for k, v in data.stats.items():
            ws3.cell(row=row, column=1, value=k)
            ws3.cell(row=row, column=2, value=str(v))
            row += 1

        wb.save(output_path)
        log.info(f"Excel report saved: {output_path}")
        return True
    except Exception as e:
        log.error(f"Excel report failed: {e}")
        return False


# ═══════════════════════════════════════════════════════════════════════════════
# PDF REPORT
# ═══════════════════════════════════════════════════════════════════════════════

def generate_pdf_report(data: ReportData, output_path: str) -> bool:
    """Генерация PDF через HTML → weasyprint"""
    try:
        from weasyprint import HTML
    except ImportError:
        log.error("weasyprint not installed. Run: pip install weasyprint")
        return False

    try:
        # Generate HTML first, then convert
        html = _HTML_TEMPLATE.format(
            title=_html.escape(data.title),
            subtitle=_html.escape(data.subtitle),
            generated_at=_html.escape(data.generated_at),
            stats_html=_render_stats_html(data.stats),
            sections_html=_render_sections_html(data.sections),
            results_html=_render_results_html(data.results),
            timeline_html=_render_timeline_html(data.timeline),
        )
        # Override dark theme for PDF
        html = html.replace("--bg: #0a0e17", "--bg: #ffffff")
        html = html.replace("--fg: #c8d6e5", "--fg: #1a1a1a")
        html = html.replace("--accent: #00ff41", "--accent: #006600")
        html = html.replace("--card: #111827", "--card: #f5f5f5")
        html = html.replace("--border: #1e3a5f", "--border: #cccccc")
        html = html.replace("background:#1a2744", "background:#003366")

        HTML(string=html).write_pdf(output_path)
        log.info(f"PDF report saved: {output_path}")
        return True
    except Exception as e:
        log.error(f"PDF report failed: {e}")
        return False


__all__ = [
    'generate_html_report', 'generate_csv_report',
    'generate_excel_report', 'generate_pdf_report',
    'ReportData', 'ReportSection'
]
